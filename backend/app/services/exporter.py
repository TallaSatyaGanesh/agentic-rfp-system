import os
from typing import Dict, Any, List, Optional
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


class ProposalExporterService:
    @staticmethod
    def _set_cell_background(cell, fill_hex: str):
        """Sets the background fill color of a table cell."""
        try:
            tc_pr = cell._tc.get_or_add_tcPr()
            for child in list(tc_pr):
                if child.tag.endswith("shd"):
                    tc_pr.remove(child)
            tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))
        except Exception:
            pass

    @staticmethod
    def _format_cell(
        cell,
        text: str,
        bold: bool = False,
        color: Optional[RGBColor] = None,
        font_size: int = 10,
        bg_hex: Optional[str] = None,
        align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT
    ):
        """Helper to format cell text, styles, and background colors."""
        if bg_hex:
            ProposalExporterService._set_cell_background(cell, bg_hex)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        run = p.add_run(text)
        run.font.bold = bold
        run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = color

    @staticmethod
    def export_to_docx(
        rfp_metadata: Dict[str, Any],
        proposal_draft: Dict[str, Any],
        compliance_matrix: List[Dict[str, Any]],
        risks: List[Dict[str, Any]],
        clarifications: List[Dict[str, Any]],
        output_path: str
    ) -> str:
        """
        Exports a professional enterprise proposal response in DOCX format,
        beginning with a dynamic Executive Compliance & Risk Summary followed
        by the complete proposal content and full formal compliance matrix appendix.
        """
        doc = docx.Document()

        # Set standard 1 inch margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # ----------------------------------------------------------------------
        # COVER / TITLE HEADER
        # ----------------------------------------------------------------------
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(proposal_draft.get("title", "PROPOSAL RESPONSE"))
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(26, 54, 93)  # Navy blue
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle_para = doc.add_paragraph()
        sub_run = subtitle_para.add_run(
            f"Submitted in Response to RFP: {rfp_metadata.get('title', 'Enterprise Tender')}\n"
            f"Issued by: {rfp_metadata.get('issuer', 'Valued Client')} | Due: {rfp_metadata.get('submission_deadline', 'Specified')}"
        )
        sub_run.font.size = Pt(12)
        sub_run.font.italic = True
        subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        divider = doc.add_paragraph()
        d_run = divider.add_run("―" * 50)
        d_run.font.color.rgb = RGBColor(203, 213, 225)
        divider.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # ----------------------------------------------------------------------
        # SECTION 1: EXECUTIVE COMPLIANCE & RISK SUMMARY
        # ----------------------------------------------------------------------
        summary_heading = doc.add_heading("1. Executive Compliance & Risk Summary", level=1)
        summary_heading.style.font.color.rgb = RGBColor(26, 54, 93)

        lead_para = doc.add_paragraph()
        lead_para.add_run(
            "This executive summary provides a high-level operational and compliance assessment of the proposal "
            "response against all extracted tender requirements, identified operational/technical risks, and proposed "
            "clarification questions."
        )

        # Compute dynamic metrics from actual compliance matrix results
        total_reqs = len(compliance_matrix)
        compliant_items = [c for c in compliance_matrix if c.get("status") == "COMPLIANT"]
        partially_compliant_items = [c for c in compliance_matrix if c.get("status") == "PARTIALLY_COMPLIANT"]
        non_compliant_items = [c for c in compliance_matrix if c.get("status") == "NON_COMPLIANT"]
        info_required_items = [
            c for c in compliance_matrix
            if c.get("status") == "INFORMATION_REQUIRED" or "INFORMATION" in c.get("status", "")
        ]

        compliant_count = len(compliant_items)
        partial_count = len(partially_compliant_items)
        non_compliant_count = len(non_compliant_items)
        info_req_count = len(info_required_items)
        compliance_pct = round((compliant_count / total_reqs * 100), 1) if total_reqs > 0 else 0.0

        # 1.1 Compliance Scorecard Table
        sub_h1 = doc.add_heading("1.1 Compliance Scorecard & Metric Breakdown", level=2)
        sub_h1.style.font.color.rgb = RGBColor(30, 64, 175)

        metric_table = doc.add_table(rows=2, cols=5)
        metric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        metric_table.style = 'Light Shading Accent 1'

        m_headers = [
            "Total Evaluated",
            "Full Compliance",
            "Partial Compliance",
            "Non-Compliant Gaps",
            "Information Needed"
        ]
        m_hdr_cells = metric_table.rows[0].cells
        for i, h_text in enumerate(m_headers):
            ProposalExporterService._format_cell(
                m_hdr_cells[i],
                h_text,
                bold=True,
                color=RGBColor(255, 255, 255),
                bg_hex="1E3A8A",
                align=WD_ALIGN_PARAGRAPH.CENTER
            )

        m_val_cells = metric_table.rows[1].cells
        ProposalExporterService._format_cell(
            m_val_cells[0],
            f"{total_reqs} Requirements",
            bold=True,
            font_size=11,
            bg_hex="F8FAFC",
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
        ProposalExporterService._format_cell(
            m_val_cells[1],
            f"{compliant_count} ({compliance_pct}%)",
            bold=True,
            font_size=11,
            color=RGBColor(22, 101, 52),
            bg_hex="DCFCE7",
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
        ProposalExporterService._format_cell(
            m_val_cells[2],
            f"{partial_count} Items",
            bold=True,
            font_size=11,
            color=RGBColor(161, 98, 7) if partial_count > 0 else RGBColor(71, 85, 105),
            bg_hex="FEF3C7" if partial_count > 0 else "F8FAFC",
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
        ProposalExporterService._format_cell(
            m_val_cells[3],
            f"{non_compliant_count} Gaps",
            bold=True,
            font_size=11,
            color=RGBColor(185, 28, 28) if non_compliant_count > 0 else RGBColor(71, 85, 105),
            bg_hex="FEE2E2" if non_compliant_count > 0 else "F8FAFC",
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
        ProposalExporterService._format_cell(
            m_val_cells[4],
            f"{info_req_count} Items",
            bold=True,
            font_size=11,
            color=RGBColor(194, 65, 12) if info_req_count > 0 else RGBColor(71, 85, 105),
            bg_hex="FFEDD5" if info_req_count > 0 else "F8FAFC",
            align=WD_ALIGN_PARAGRAPH.CENTER
        )

        doc.add_paragraph()

        # 1.2 Non-Compliant Requirements Summary
        sub_h2 = doc.add_heading("1.2 Non-Compliant Requirements & Capability Gaps", level=2)
        sub_h2.style.font.color.rgb = RGBColor(30, 64, 175)

        if non_compliant_count > 0:
            nc_table = doc.add_table(rows=1, cols=4)
            nc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            nc_table.style = 'Light Shading Accent 1'

            nc_headers = ["Req ID", "Category", "Requirement Description", "Gap Description / Rationale"]
            nc_hdr_cells = nc_table.rows[0].cells
            for i, h_text in enumerate(nc_headers):
                ProposalExporterService._format_cell(
                    nc_hdr_cells[i],
                    h_text,
                    bold=True,
                    color=RGBColor(255, 255, 255),
                    bg_hex="991B1B",
                    align=WD_ALIGN_PARAGRAPH.LEFT
                )

            for item in non_compliant_items:
                row_cells = nc_table.add_row().cells
                req_code = item.get("req_code", "N/A")
                category = item.get("category", "General")
                text = item.get("requirement_text") or item.get("text") or "Requirement text not specified."
                reason = item.get("notes") or item.get("evidence_text") or item.get("company_source_doc") or "Out of platform scope."

                ProposalExporterService._format_cell(row_cells[0], req_code, bold=True, color=RGBColor(185, 28, 28))
                ProposalExporterService._format_cell(row_cells[1], category)
                ProposalExporterService._format_cell(row_cells[2], text)
                ProposalExporterService._format_cell(row_cells[3], reason)
        else:
            p_nc = doc.add_paragraph()
            r_nc = p_nc.add_run("✓ Zero non-compliant requirements identified. The proposed solution meets or exceeds all core baseline specifications.")
            r_nc.font.color.rgb = RGBColor(22, 101, 52)
            r_nc.font.bold = True

        doc.add_paragraph()

        # 1.3 Partially Compliant Requirements Summary
        sub_h3 = doc.add_heading("1.3 Partially Compliant Requirements & Caveats", level=2)
        sub_h3.style.font.color.rgb = RGBColor(30, 64, 175)

        if partial_count > 0:
            pc_table = doc.add_table(rows=1, cols=4)
            pc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            pc_table.style = 'Light Shading Accent 1'

            pc_headers = ["Req ID", "Category", "Requirement Description", "Compliance Caveat / Customization Path"]
            pc_hdr_cells = pc_table.rows[0].cells
            for i, h_text in enumerate(pc_headers):
                ProposalExporterService._format_cell(
                    pc_hdr_cells[i],
                    h_text,
                    bold=True,
                    color=RGBColor(255, 255, 255),
                    bg_hex="92400E",
                    align=WD_ALIGN_PARAGRAPH.LEFT
                )

            for item in partially_compliant_items:
                row_cells = pc_table.add_row().cells
                req_code = item.get("req_code", "N/A")
                category = item.get("category", "General")
                text = item.get("requirement_text") or item.get("text") or "Requirement text not specified."
                notes = item.get("notes") or item.get("evidence_text") or item.get("company_source_doc") or "Supported with standard configuration or workflow extension."

                ProposalExporterService._format_cell(row_cells[0], req_code, bold=True, color=RGBColor(161, 98, 7))
                ProposalExporterService._format_cell(row_cells[1], category)
                ProposalExporterService._format_cell(row_cells[2], text)
                ProposalExporterService._format_cell(row_cells[3], notes)
        else:
            p_pc = doc.add_paragraph()
            r_pc = p_pc.add_run("✓ No partial compliance caveats recorded. All fulfilled criteria comply fully with native capabilities.")
            r_pc.font.color.rgb = RGBColor(22, 101, 52)
            r_pc.font.bold = True

        doc.add_paragraph()

        # 1.4 Information-Required Action Items
        sub_h4 = doc.add_heading("1.4 Information-Required Action Items", level=2)
        sub_h4.style.font.color.rgb = RGBColor(30, 64, 175)

        if info_req_count > 0:
            ir_table = doc.add_table(rows=1, cols=3)
            ir_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ir_table.style = 'Light Shading Accent 1'

            ir_headers = ["Req ID", "Category", "Information Needed / Next Step"]
            ir_hdr_cells = ir_table.rows[0].cells
            for i, h_text in enumerate(ir_headers):
                ProposalExporterService._format_cell(
                    ir_hdr_cells[i],
                    h_text,
                    bold=True,
                    color=RGBColor(255, 255, 255),
                    bg_hex="C2410C",
                    align=WD_ALIGN_PARAGRAPH.LEFT
                )

            for item in info_required_items:
                row_cells = ir_table.add_row().cells
                req_code = item.get("req_code", "N/A")
                category = item.get("category", "General")
                action = item.get("notes") or item.get("evidence_text") or "Detailed technical clarification requested from issuer/vendor."

                ProposalExporterService._format_cell(row_cells[0], req_code, bold=True, color=RGBColor(194, 65, 12))
                ProposalExporterService._format_cell(row_cells[1], category)
                ProposalExporterService._format_cell(row_cells[2], action)
        else:
            p_ir = doc.add_paragraph()
            r_ir = p_ir.add_run("✓ All requirement specifications are fully clear. No outstanding information requests remain pending.")
            r_ir.font.color.rgb = RGBColor(22, 101, 52)
            r_ir.font.bold = True

        doc.add_paragraph()

        # 1.5 Key Risks & Mitigations Summary
        sub_h5 = doc.add_heading("1.5 Key Risk Register & Mitigation Strategy", level=2)
        sub_h5.style.font.color.rgb = RGBColor(30, 64, 175)

        if risks:
            risk_table = doc.add_table(rows=1, cols=4)
            risk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            risk_table.style = 'Light Shading Accent 1'

            risk_headers = ["Category", "Severity", "Risk Description", "Mitigation Strategy"]
            risk_hdr_cells = risk_table.rows[0].cells
            for i, h_text in enumerate(risk_headers):
                ProposalExporterService._format_cell(
                    risk_hdr_cells[i],
                    h_text,
                    bold=True,
                    color=RGBColor(255, 255, 255),
                    bg_hex="374151",
                    align=WD_ALIGN_PARAGRAPH.LEFT
                )

            for r in risks:
                row_cells = risk_table.add_row().cells
                cat = r.get("category", "Operational")
                sev = r.get("severity", "Medium")
                desc = r.get("description", "Identified engagement risk.")
                mit = r.get("mitigation_strategy", "Standard mitigation protocol.")

                ProposalExporterService._format_cell(row_cells[0], cat, bold=True)
                
                # Severity formatting
                sev_upper = str(sev).upper()
                sev_color = RGBColor(185, 28, 28) if "HIGH" in sev_upper or "CRITICAL" in sev_upper else (
                    RGBColor(161, 98, 7) if "MEDIUM" in sev_upper else RGBColor(22, 101, 52)
                )
                ProposalExporterService._format_cell(row_cells[1], sev, bold=True, color=sev_color)
                ProposalExporterService._format_cell(row_cells[2], desc)
                ProposalExporterService._format_cell(row_cells[3], mit)
        else:
            p_risk = doc.add_paragraph()
            r_risk = p_risk.add_run("✓ No high or critical risks identified for this proposal engagement.")
            r_risk.font.color.rgb = RGBColor(22, 101, 52)
            r_risk.font.bold = True

        doc.add_paragraph()

        # 1.6 Tender Clarification Questions
        sub_h6 = doc.add_heading("1.6 Tender Clarification Questions", level=2)
        sub_h6.style.font.color.rgb = RGBColor(30, 64, 175)

        if clarifications:
            clarif_table = doc.add_table(rows=1, cols=4)
            clarif_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            clarif_table.style = 'Light Shading Accent 1'

            clarif_headers = ["Q#", "RFP Section Ref", "Question for Issuer", "Technical Rationale"]
            clarif_hdr_cells = clarif_table.rows[0].cells
            for i, h_text in enumerate(clarif_headers):
                ProposalExporterService._format_cell(
                    clarif_hdr_cells[i],
                    h_text,
                    bold=True,
                    color=RGBColor(255, 255, 255),
                    bg_hex="1F2937",
                    align=WD_ALIGN_PARAGRAPH.LEFT
                )

            for idx, q in enumerate(clarifications, 1):
                row_cells = clarif_table.add_row().cells
                q_num = f"Q-{q.get('q_number', idx)}"
                sec_ref = q.get("rfp_section_reference", "General")
                q_text = q.get("question_text", "Clarification requested.")
                rationale = q.get("rationale") or "To ensure precision in solution architecture and delivery milestones."

                ProposalExporterService._format_cell(row_cells[0], q_num, bold=True, color=RGBColor(30, 64, 175))
                ProposalExporterService._format_cell(row_cells[1], sec_ref)
                ProposalExporterService._format_cell(row_cells[2], q_text)
                ProposalExporterService._format_cell(row_cells[3], rationale)
        else:
            p_cl = doc.add_paragraph()
            p_cl.add_run("No formal clarification questions submitted for this response.")

        # ----------------------------------------------------------------------
        # SECTION 2: DETAILED TECHNICAL & COMMERCIAL PROPOSAL CONTENT
        # ----------------------------------------------------------------------
        doc.add_page_break()
        prop_h = doc.add_heading("2. Detailed Proposal Response", level=1)
        prop_h.style.font.color.rgb = RGBColor(26, 54, 93)

        sections = proposal_draft.get("sections", [])
        if sections:
            for sec in sections:
                h = doc.add_heading(sec.get("section_title", "Section"), level=2)
                h.style.font.color.rgb = RGBColor(30, 64, 175)

                content = sec.get("content_markdown", "")
                paragraphs = content.split("\n\n")
                for p_text in paragraphs:
                    p_text_clean = p_text.strip()
                    if not p_text_clean:
                        continue

                    # Special styling for INFORMATION REQUIRED callouts
                    if "INFORMATION REQUIRED" in p_text_clean:
                        callout = doc.add_paragraph()
                        c_run = callout.add_run("⚠️ " + p_text_clean.replace(">", "").strip())
                        c_run.font.bold = True
                        c_run.font.color.rgb = RGBColor(180, 83, 9)  # Amber-700
                    elif p_text_clean.startswith("### "):
                        doc.add_heading(p_text_clean.replace("### ", ""), level=3)
                    elif p_text_clean.startswith("#### "):
                        doc.add_heading(p_text_clean.replace("#### ", ""), level=4)
                    elif p_text_clean.startswith("- "):
                        items = p_text_clean.split("\n")
                        for it in items:
                            if it.startswith("- "):
                                doc.add_paragraph(it.replace("- ", ""), style='List Bullet')
                    else:
                        doc.add_paragraph(p_text_clean)
        else:
            # Fallback to full markdown text
            for line in proposal_draft.get("full_markdown", "").split("\n"):
                doc.add_paragraph(line)

        # ----------------------------------------------------------------------
        # APPENDIX A: COMPLETE FORMAL REQUIREMENT COMPLIANCE MATRIX
        # ----------------------------------------------------------------------
        doc.add_page_break()
        mat_heading = doc.add_heading("Appendix A: Formal Requirement Compliance Matrix", level=1)
        mat_heading.style.font.color.rgb = RGBColor(26, 54, 93)

        if compliance_matrix:
            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Light Shading Accent 1'

            hdr_cells = table.rows[0].cells
            headers = ["Req ID", "Category", "Compliance Verdict", "Source Evidence / Action Required"]
            for i, header_text in enumerate(headers):
                ProposalExporterService._format_cell(
                    hdr_cells[i],
                    header_text,
                    bold=True,
                    color=RGBColor(255, 255, 255),
                    bg_hex="1E3A8A",
                    align=WD_ALIGN_PARAGRAPH.LEFT
                )

            for item in compliance_matrix:
                row_cells = table.add_row().cells
                row_cells[0].text = item.get("req_code", "")
                row_cells[1].text = item.get("category", "")

                # Color coded verdict
                status = item.get("status", "")
                row_cells[2].text = status
                verdict_run = row_cells[2].paragraphs[0].runs[0]
                verdict_run.font.bold = True
                if status == "COMPLIANT":
                    verdict_run.font.color.rgb = RGBColor(22, 101, 52)
                elif status == "PARTIALLY_COMPLIANT":
                    verdict_run.font.color.rgb = RGBColor(161, 98, 7)
                elif status == "NON_COMPLIANT":
                    verdict_run.font.color.rgb = RGBColor(185, 28, 28)
                else:
                    verdict_run.font.color.rgb = RGBColor(194, 65, 12)

                evidence = item.get("company_source_doc") or item.get("notes") or item.get("evidence_text") or "Evidence Required"
                row_cells[3].text = evidence

        # Ensure directory exists and save
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path
